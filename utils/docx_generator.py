"""
Optional DOCX generation for offer letters.
"""
import io
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


TERICSOFT_BLUE = RGBColor(0x1A, 0x3C, 0x6B)


def _add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(11) if level == 1 else Pt(10)
    run.font.color.rgb = TERICSOFT_BLUE
    return p


def _add_paragraph(doc, text, bold=False, italic=False, size=10):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    return p


def _add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    run.font.size = Pt(10)
    return p


def generate_internship_docx(data: dict) -> bytes:
    doc = Document()

    # Margins
    for section in doc.sections:
        section.top_margin = Cm(3.5)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # Date
    _add_paragraph(doc, data.get("letter_date", ""), bold=True)

    # Subject
    p = doc.add_paragraph()
    run = p.add_run("Subject: Internship at Tericsoft Technology")
    run.bold = True
    run.underline = True
    run.font.size = Pt(10)

    # Salutation
    _add_paragraph(doc, f"Dear {data.get('name', 'Candidate')},")

    # Body
    role = data.get("role", "Intern")
    duration = data.get("duration", "3 months")
    joining_date = data.get("joining_date", "")
    _add_paragraph(
        doc,
        f"In reference to your application, we would like to congratulate you on your internship "
        f"for the position of {role} for {duration}. Your internship is scheduled to start effective "
        f"from {joining_date}. All of us at Tericsoft are excited that you will be joining our team."
    )

    responsibilities = data.get("responsibilities", "")
    if responsibilities:
        _add_paragraph(doc, responsibilities)

    _add_paragraph(doc, "The in-depth details of the internship will be shared by your mentor.")

    _add_paragraph(doc, f"Date of Joining: {joining_date}", bold=True)
    _add_paragraph(doc, f"Timings: {data.get('timings', '10:30am – 7pm')}", bold=True)
    _add_paragraph(doc, "Weekdays: Monday to Friday", bold=True)

    doc.add_paragraph()
    _add_paragraph(doc, "Again, congratulations and we look forward to working with you.")

    # Signature
    doc.add_paragraph()
    _add_paragraph(doc, "Yours sincerely,")
    doc.add_paragraph()
    doc.add_paragraph()
    _add_paragraph(doc, "Abdul Rahman", bold=True)
    _add_paragraph(doc, "Director")
    _add_paragraph(doc, "Tericsoft Technology")

    # Acceptance
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("_" * 35 + "  Date - " + "_" * 35).font.size = Pt(10)
    _add_paragraph(doc, data.get("full_name", ""))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generate_fulltime_docx(data: dict, comp_table_data: list | None = None) -> bytes:
    doc = Document()

    for section in doc.sections:
        section.top_margin = Cm(3.5)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2.5)
        section.right_margin = Cm(2.5)

    # Date & recipient
    _add_paragraph(doc, data.get("letter_date", ""), bold=True)
    _add_paragraph(doc, data.get("full_name", ""))
    _add_paragraph(doc, data.get("email", ""))

    # Salutation
    _add_paragraph(doc, f"Dear {data.get('name', 'Candidate')},")

    role = data.get("role", "Executive")
    reporting_to = data.get("reporting_to", "the Director")
    probation = data.get("probation_period", "3")
    joining_date = data.get("joining_date", "")

    _add_paragraph(
        doc,
        f"It gives us great pleasure to extend you an offer to join the Tericsoft Technology team! "
        f"We would like you to join in the role of {role} reporting to {reporting_to}. "
        f"We were impressed with your strong ability combined with your enthusiasm to support the "
        f"development of Tericsoft's vision & mission. Together, we feel that these attributes will make "
        f"you an outstanding fit as part of Tericsoft's technology team."
    )
    _add_paragraph(doc, "This letter contains the relevant information regarding your offer.")

    # Term
    _add_heading(doc, "1.  Term:", level=2)
    _add_paragraph(
        doc,
        f"As {role}, you will start on a {probation}-month probation period with the intention "
        f"of extending into the permanent role. Your services will be confirmed in writing after the "
        f"successful completion of your probation period. The probation period may be extended if your "
        f"performance does not meet expectations."
    )

    # Responsibilities
    _add_heading(doc, "2.  Responsibilities:", level=2)
    _add_paragraph(doc, "This role covers several key areas:", italic=True)

    responsibilities_text = data.get("responsibilities", "")
    for line in responsibilities_text.split("\n"):
        line = line.strip()
        if not line:
            continue
        if line.endswith(":") and not line.startswith("•"):
            _add_paragraph(doc, line, bold=True)
        elif line.startswith("•") or line.startswith("-"):
            _add_bullet(doc, line.lstrip("•- ").strip())
        else:
            _add_paragraph(doc, line)

    # Compensation
    _add_heading(doc, "Compensation:", level=2)
    monthly = data.get("monthly_salary", "")
    ctc = data.get("ctc", "")
    _add_paragraph(
        doc,
        f"Salary: You will receive a fixed salary of INR {monthly}/- per month "
        f"(Inclusive of all taxes & benefits). Your fixed CTC (Cost to Company) will be INR {ctc}/- per annum."
    )

    # Compensation table
    if comp_table_data:
        headers = ["Component", "Percentage", "Per Month", "Per Annum"]
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        hdr_cells = table.rows[0].cells
        for i, h in enumerate(headers):
            hdr_cells[i].text = h
            hdr_cells[i].paragraphs[0].runs[0].bold = True
        for row in comp_table_data:
            row_cells = table.add_row().cells
            for i, val in enumerate(row):
                row_cells[i].text = str(val)

    # Closing clauses
    _add_paragraph(doc, "Offer stands cancelled in case of any deviations in information/if not reported before the date of acceptance.")
    _add_paragraph(doc, "Please reply with your confirmation of acceptance of offer. You will have to submit certain documents as a part of the onboarding process on the joining date.")
    _add_paragraph(
        doc,
        f"Notice Period – During probation period, the notice period stands for 30 days. Post probation, "
        f"the notice period stands for 60 days. You will be on probation for {probation} months from the date of your joining."
    )
    _add_paragraph(doc, f"Date of Joining: {joining_date}.", bold=True)
    _add_paragraph(doc, "I am confident you will have a rewarding experience with a phenomenal team. We look forward to welcoming you to Tericsoft Technology.")

    # Signature
    doc.add_paragraph()
    _add_paragraph(doc, "Sincerely,")
    doc.add_paragraph()
    doc.add_paragraph()
    _add_paragraph(doc, "Abdul Rahman", bold=True)
    _add_paragraph(doc, "Director")
    _add_paragraph(doc, "Tericsoft Technology")

    # Acceptance
    doc.add_paragraph()
    p = doc.add_paragraph()
    p.add_run("_" * 35 + "  Date - " + "_" * 35).font.size = Pt(10)
    _add_paragraph(doc, data.get("full_name", ""))

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
